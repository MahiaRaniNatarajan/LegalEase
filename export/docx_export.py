"""Member 4: DOCX export."""


from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def export_docx(content: str, path: str) -> str:
    doc = Document()

    # Document title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run("LEGALEASE")
    run.bold = True
    run.font.size = Pt(20)

    # Document content
    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()

        if not paragraph:
            continue

        # Treat numbered sections as headings
        if paragraph[0].isdigit() and "." in paragraph[:4]:
            heading = doc.add_paragraph()
            run = heading.add_run(paragraph)
            run.bold = True
            run.font.size = Pt(13)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)

            run = p.add_run(paragraph)
            run.font.size = Pt(11)

    doc.save(path)

    return path
